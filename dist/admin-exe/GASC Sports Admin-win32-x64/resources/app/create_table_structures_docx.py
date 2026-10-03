import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/></w:tblBorders>')
    tblPr.append(borders)

def create_appendix_d_docx():
    doc = docx.Document()
    
    # Page Setup - Margins 1 inch
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("APPENDIX D: KEY TABLE STRUCTURES (SQL)")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78) # Dark Navy
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Government Arts and Science College, Idappadi\nDepartment of Computer Science & Physical Education")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    
    doc.add_paragraph() # Spacer
    
    tables_data = [
        {
            "id": "D.1",
            "name": "users Table",
            "desc": "Stores core account information, bonafide register numbers, role permissions, academic department details, and contact profiles for students and administrative staff.",
            "sql": """CREATE TABLE IF NOT EXISTS users (
    id               UUID PRIMARY KEY DEFAULT uuid_generate_v4(),
    name             VARCHAR(255) NOT NULL,
    register_number  VARCHAR(100) UNIQUE,
    email            VARCHAR(255) UNIQUE NOT NULL,
    password         VARCHAR(255) NOT NULL,
    role             VARCHAR(50)  DEFAULT 'student'
                     CHECK (role IN ('admin', 'student')),
    department       VARCHAR(255) DEFAULT 'General',
    year             VARCHAR(50)  DEFAULT 'I Year'
                     CHECK (year IN ('I Year','II Year','III Year','Faculty','Other')),
    section          VARCHAR(10)  DEFAULT 'A',
    gender           VARCHAR(20)  DEFAULT 'Male'
                     CHECK (gender IN ('Male','Female','Other')),
    dob              DATE,
    mobile           VARCHAR(50),
    profile_photo    TEXT         DEFAULT '/images/default-avatar.png',
    status           VARCHAR(50)  DEFAULT 'Active'
                     CHECK (status IN ('Active','Inactive','Suspended')),
    created_at       TIMESTAMPTZ  DEFAULT NOW(),
    updated_at       TIMESTAMPTZ  DEFAULT NOW()
);""",
            "columns": [
                ("id", "UUID", "PRIMARY KEY", "Unique user ID (uuid_generate_v4())"),
                ("name", "VARCHAR(255)", "NOT NULL", "Full student / staff member name"),
                ("register_number", "VARCHAR(100)", "UNIQUE", "College Roll / Register number"),
                ("email", "VARCHAR(255)", "UNIQUE, NOT NULL", "Email address for OTP verification"),
                ("password", "VARCHAR(255)", "NOT NULL", "Encrypted password hash"),
                ("role", "VARCHAR(50)", "CHECK ('admin','student')", "System role (Default: 'student')"),
                ("department", "VARCHAR(255)", "DEFAULT 'General'", "Academic department name"),
                ("year", "VARCHAR(50)", "CHECK ('I Year','II Year',...)", "Class year (Default: 'I Year')"),
                ("section", "VARCHAR(10)", "DEFAULT 'A'", "Class section"),
                ("gender", "VARCHAR(20)", "CHECK ('Male','Female','Other')", "Gender identity (Default: 'Male')"),
                ("dob", "DATE", "NULL", "Date of birth"),
                ("mobile", "VARCHAR(50)", "NULL", "Contact phone number"),
                ("profile_photo", "TEXT", "DEFAULT '/images/...'", "Profile photo URL path"),
                ("status", "VARCHAR(50)", "CHECK ('Active','Inactive',...)", "Account status (Default: 'Active')"),
                ("created_at", "TIMESTAMPTZ", "DEFAULT NOW()", "Record creation timestamp"),
                ("updated_at", "TIMESTAMPTZ", "DEFAULT NOW()", "Record update timestamp")
            ]
        },
        {
            "id": "D.2",
            "name": "equipment_transactions Table",
            "desc": "Tracks digital borrowing, return, damage assessment, fines, and stock movement of collegiate sports gear.",
            "sql": """CREATE TABLE IF NOT EXISTS equipment_transactions (
    id                   UUID PRIMARY KEY DEFAULT uuid_generate_v4(),
    student_id           UUID NOT NULL REFERENCES users(id) ON DELETE CASCADE,
    student_name         VARCHAR(255),
    register_number      VARCHAR(100),
    equipment_id         UUID NOT NULL REFERENCES equipment(id) ON DELETE CASCADE,
    equipment_name       VARCHAR(255),
    quantity             INTEGER  NOT NULL DEFAULT 1,
    issue_date           TIMESTAMPTZ DEFAULT NOW(),
    expected_return_date TIMESTAMPTZ NOT NULL,
    return_date          TIMESTAMPTZ,
    status               VARCHAR(50)  DEFAULT 'Issued'
                         CHECK (status IN ('Issued','Returned','Damaged','Lost','Overdue')),
    return_condition     VARCHAR(100) DEFAULT 'Pending'
                         CHECK (return_condition IN
                               ('Good','Damaged','Lost','Partially Damaged','Pending')),
    damage_description   TEXT,
    fine_amount          NUMERIC(10,2) DEFAULT 0.00,
    purpose              VARCHAR(255) DEFAULT 'College Practice / Match',
    issued_by            VARCHAR(255) DEFAULT 'Sports Incharge',
    created_at           TIMESTAMPTZ DEFAULT NOW()
);""",
            "columns": [
                ("id", "UUID", "PRIMARY KEY", "Unique transaction record ID"),
                ("student_id", "UUID", "FK -> users(id)", "Foreign key to borrowing student"),
                ("student_name", "VARCHAR(255)", "NULL", "Student name snapshot"),
                ("register_number", "VARCHAR(100)", "NULL", "Student register number snapshot"),
                ("equipment_id", "UUID", "FK -> equipment(id)", "Foreign key to sports equipment"),
                ("equipment_name", "VARCHAR(255)", "NULL", "Equipment item name"),
                ("quantity", "INTEGER", "NOT NULL, DEFAULT 1", "Issued gear quantity"),
                ("issue_date", "TIMESTAMPTZ", "DEFAULT NOW()", "Checkout timestamp"),
                ("expected_return_date", "TIMESTAMPTZ", "NOT NULL", "Scheduled return deadline"),
                ("return_date", "TIMESTAMPTZ", "NULL", "Actual check-in timestamp"),
                ("status", "VARCHAR(50)", "CHECK ('Issued','Returned',...)", "Issue status (Default: 'Issued')"),
                ("return_condition", "VARCHAR(100)", "CHECK ('Good','Damaged',...)", "Condition upon check-in"),
                ("damage_description", "TEXT", "NULL", "Details of damage if applicable"),
                ("fine_amount", "NUMERIC(10,2)", "DEFAULT 0.00", "Penalty fee for late/damaged item"),
                ("purpose", "VARCHAR(255)", "DEFAULT 'College Practice...'", "Reason for issue"),
                ("issued_by", "VARCHAR(255)", "DEFAULT 'Sports Incharge'", "Staff issuer identifier"),
                ("created_at", "TIMESTAMPTZ", "DEFAULT NOW()", "Record creation timestamp")
            ]
        },
        {
            "id": "D.3",
            "name": "attendance Table",
            "desc": "Logs student athlete attendance during daily coaching and morning/evening sports training sessions.",
            "sql": """CREATE TABLE IF NOT EXISTS attendance (
    id                  UUID PRIMARY KEY DEFAULT uuid_generate_v4(),
    practice_session_id UUID NOT NULL
                        REFERENCES practice_sessions(id) ON DELETE CASCADE,
    student_id          UUID NOT NULL
                        REFERENCES users(id) ON DELETE CASCADE,
    status              VARCHAR(50) DEFAULT 'Present'
                        CHECK (status IN ('Present','Absent','Late','Excused')),
    date                DATE NOT NULL,
    remarks             TEXT,
    recorded_by         VARCHAR(255) DEFAULT 'Sports Incharge',
    created_at          TIMESTAMPTZ DEFAULT NOW(),
    UNIQUE (practice_session_id, student_id)  -- Prevents duplicate attendance
);""",
            "columns": [
                ("id", "UUID", "PRIMARY KEY", "Unique attendance entry ID"),
                ("practice_session_id", "UUID", "FK -> practice_sessions(id)", "Foreign key to practice session"),
                ("student_id", "UUID", "FK -> users(id)", "Foreign key to student athlete"),
                ("status", "VARCHAR(50)", "CHECK ('Present','Absent',...)", "Attendance status (Default: 'Present')"),
                ("date", "DATE", "NOT NULL", "Practice session date"),
                ("remarks", "TEXT", "NULL", "Performance or absence notes"),
                ("recorded_by", "VARCHAR(255)", "DEFAULT 'Sports Incharge'", "Physical Director / Coach recorder"),
                ("created_at", "TIMESTAMPTZ", "DEFAULT NOW()", "Record creation timestamp"),
                ("UNIQUE Constraint", "COMPOSITE", "UNIQUE(session_id, student_id)", "Prevents duplicate entry per session")
            ]
        }
    ]

    for item in tables_data:
        # Section Heading
        p_h = doc.add_paragraph()
        run_h = p_h.add_run(f"{item['id']} – {item['name']}")
        run_h.font.name = "Arial"
        run_h.font.size = Pt(14)
        run_h.font.bold = True
        run_h.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
        
        # Description
        p_d = doc.add_paragraph()
        run_d = p_d.add_run(f"Description: {item['desc']}")
        run_d.font.name = "Calibri"
        run_d.font.size = Pt(11)
        run_d.font.italic = True
        
        # SQL Code block
        p_sql_title = doc.add_paragraph()
        run_sqlt = p_sql_title.add_run("SQL Schema Definition:")
        run_sqlt.font.name = "Arial"
        run_sqlt.font.size = Pt(10)
        run_sqlt.font.bold = True
        
        # SQL Box (Table with 1 cell shading)
        sql_tbl = doc.add_table(rows=1, cols=1)
        sql_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = sql_tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F2F2F2")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        
        p_code = cell.paragraphs[0]
        run_code = p_code.add_run(item['sql'])
        run_code.font.name = "Consolas"
        run_code.font.size = Pt(9.5)
        run_code.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
        
        doc.add_paragraph() # Spacer
        
        # Data Dictionary Table
        p_tbl_title = doc.add_paragraph()
        run_tblt = p_tbl_title.add_run(f"Table Structure / Data Dictionary for {item['name']}:")
        run_tblt.font.name = "Arial"
        run_tblt.font.size = Pt(10)
        run_tblt.font.bold = True
        
        table = doc.add_table(rows=len(item['columns']) + 1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="D3D3D3", sz="4")
        
        headers = ["Field Name", "Data Type", "Constraints / Key", "Description"]
        col_widths = [Inches(1.5), Inches(1.2), Inches(1.6), Inches(2.2)]
        
        # Format Header Row
        hdr_cells = table.rows[0].cells
        for i, header_text in enumerate(headers):
            hdr_cells[i].width = col_widths[i]
            set_cell_background(hdr_cells[i], "1F4E78") # Dark Navy Header
            set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
            p = hdr_cells[i].paragraphs[0]
            run = p.add_run(header_text)
            run.font.name = "Arial"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) # White text
            
        # Format Data Rows
        for row_idx, col_data in enumerate(item['columns']):
            row_cells = table.rows[row_idx + 1].cells
            bg_color = "F9FBFD" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, text in enumerate(col_data):
                row_cells[col_idx].width = col_widths[col_idx]
                set_cell_background(row_cells[col_idx], bg_color)
                set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=120, right=120)
                p = row_cells[col_idx].paragraphs[0]
                run = p.add_run(text)
                run.font.name = "Consolas" if col_idx < 2 else "Calibri"
                run.font.size = Pt(9.5) if col_idx < 2 else Pt(10)
                
        doc.add_paragraph() # Spacer between tables
        doc.add_paragraph()

    # Save Word Document
    out_c_drive = r"C:\DFD_Images\Appendix_D_Table_Structures.docx"
    out_project = r"c:\Users\ELCOT\.gemini\antigravity-ide\scratch\gasc-idappadi-sports\Appendix_D_Table_Structures.docx"
    
    doc.save(out_c_drive)
    doc.save(out_project)
    print("SUCCESSFULLY_CREATED_WORD_DOCUMENT")

create_appendix_d_docx()
